"""The Excel sizing model agrees with the Python sizing at the default inputs (values cached by Excel on save)."""

from __future__ import annotations

import warnings

import openpyxl
import pytest

from ff import sizing
from ff.sources import ROOT

XLSX = ROOT / "excel" / "Field-Force-Sizing-Model.xlsx"


@pytest.fixture(scope="module")
def wb():
    if not XLSX.exists():
        pytest.skip("workbook not built")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(XLSX, data_only=True)


def test_sheets_and_names_exist(wb):
    assert wb.sheetnames == ["Read me", "Inputs", "Sizing", "States", "Checks", "Data", "State facts"]
    for name in ("Franchise", "Capacity", "seg_calls", "seg_share", "d_work", "s_patients"):
        assert name in wb.defined_names


def test_default_diabetes_reps_equal_the_python_sizing(wb):
    fr, rules, segs = sizing.load_franchises()["diabetes"], sizing.load_rules(), sizing.load_county_segments()
    py = sizing.reps_needed(sum(sizing.county_workload("diabetes", segs, rules).values()), fr)
    assert wb["Inputs"]["C4"].value == "diabetes"
    assert wb["Sizing"]["C4"].value == pytest.approx(py, rel=1e-9)
    assert wb["Checks"]["E12"].value == "ALL OK"


def test_states_sheet_has_51_states_and_a_total_that_matches_the_headline(wb):
    ws = wb["States"]
    assert [ws.cell(r, 1).value for r in (6, 56)] != [None, None]
    assert ws["D57"].value == pytest.approx(wb["Sizing"]["C4"].value, rel=1e-12)
    assert sum(1 for r in range(6, 57) if ws.cell(r, 1).value) == 51


def test_workbook_has_live_formulas_not_pasted_results(wb):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        f = openpyxl.load_workbook(XLSX)
    assert str(f["Sizing"]["C4"].value).startswith("=")
    assert str(f["States"]["C6"].value).startswith("=SUMIFS")
    assert str(f["Data"]["H2"].value).startswith("=")
